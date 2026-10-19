"""Export Excel (.xlsx) : classeur multi-feuilles."""

from __future__ import annotations

import io

from openpyxl import load_workbook

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_excel_requires_session(client):
    assert client.get("/api/admin/export.xlsx").status_code == 401


def test_excel_export_has_expected_sheets(admin_client, create_item):
    create_item(
        name="Rafale M",
        item_type="DIECAST",
        country="France",
        estimated_value=120,
        purchase_price=100,
        purchase_date="2024-05-01",
        diecast={"scale": "1/72", "manufacturer": "Hobby Master"},
    )
    create_item(name="Souhait", status="WANTED", estimated_value=80)
    create_item(
        name="Precommande",
        status="PREORDER",
        preorder_release_date="2027-01-01",
        preorder_shop="Boutique",
        preorder_deposit=20,
    )

    response = admin_client.get("/api/admin/export.xlsx")

    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX_MIME
    assert response.content[:2] == b"PK"  # archive ZIP (xlsx)

    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Collection", "Souhaits", "Précommandes", "Synthèse"]

    collection = workbook["Collection"]
    headers = [cell.value for cell in collection[1]]
    assert headers[:6] == ["ID", "Nom", "Type", "Statut", "Pays", "Modèle"]
    assert "Classe" in headers
    assert "Valeur estimée (EUR)" in headers
    assert "Plus-value (EUR)" in headers
    # Seuls les objets possédés figurent dans la feuille Collection.
    names = [row[1] for row in collection.iter_rows(min_row=2, values_only=True)]
    assert names == ["Rafale M"]
    assert collection.cell(row=2, column=3).value == "Maquette"
    assert collection.cell(row=2, column=10).value.year == 2024  # date d'achat

    wishes = workbook["Souhaits"]
    assert [row[1] for row in wishes.iter_rows(min_row=2, values_only=True)] == ["Souhait"]

    preorders = workbook["Précommandes"]
    assert [row[1] for row in preorders.iter_rows(min_row=2, values_only=True)] == ["Precommande"]
    preorder_headers = [cell.value for cell in preorders[1]]
    assert "Sortie prévue" in preorder_headers
    assert "Boutique / fournisseur" in preorder_headers
    assert "Acompte versé" in preorder_headers

    summary = workbook["Synthèse"]
    assert summary["A1"].value == "Synthèse de la collection"


def test_excel_export_on_empty_collection(admin_client):
    response = admin_client.get("/api/admin/export.xlsx")
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook["Collection"].max_row >= 1
