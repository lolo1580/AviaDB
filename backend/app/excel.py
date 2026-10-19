"""Export Excel (.xlsx) — lisible et analysable dans un tableur.

Un classeur, quatre feuilles :
- **Collection** : objets possédés (hors souhaits et précommandes) ;
- **Souhaits** : objets recherchés, avec budget et dernier relevé ;
- **Précommandes** : objets commandés mais non reçus ;
- **Synthèse** : totaux et répartitions.

Le CSV technique (`/api/admin/export.csv`) reste le format d'import/aller-retour.
"""

from __future__ import annotations

import io
from collections.abc import Callable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from . import crud, currency, models
from .stats import compute_stats

TYPE_LABEL = {"DIECAST": "Maquette", "PATCH": "Patch", "OTHER": "Autre"}
STATUS_LABEL = {
    "EXHIBITED": "Exposé",
    "STORED_BOX": "En boîte",
    "WANTED": "Recherché",
    "PREORDER": "Précommande",
    "TRANSIT": "En transit",
}
MATERIAL_LABEL = {
    "EMBROIDERED": "Broderie",
    "PVC": "PVC",
    "LEATHER": "Cuir",
    "WOVEN": "Tissé",
}
ATTACHMENT_LABEL = {
    "HOOK_VELCRO": "Velcro crochet",
    "LOOP_VELCRO": "Velcro boucle",
    "SEW_ON": "À coudre",
}
SHAPE_LABEL = {"ROUND": "Rond", "SHIELD": "Écusson", "TAB": "Fanion", "CUT_OUT": "Découpé"}

MONEY_FORMAT = "#,##0.00"
DATE_FORMAT = "DD/MM/YYYY"

_HEADER_FILL = PatternFill("solid", fgColor="0F172A")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_TITLE_FONT = Font(bold=True, size=14, color="0F172A")
_BLOCK_FONT = Font(bold=True, size=11, color="0F172A")
_NOTE_FONT = Font(italic=True, color="64748B")

# (libellé, extracteur, format) — format ∈ {None, "money", "date", "wrap"}.
Column = tuple[str, Callable[[models.Item], object], str | None]


# --------------------------------------------------------------------------- #
# Données
# --------------------------------------------------------------------------- #
def _items(db: Session) -> list[models.Item]:
    return list(
        db.scalars(
            select(models.Item)
            .options(
                selectinload(models.Item.diecast),
                selectinload(models.Item.patch),
                selectinload(models.Item.photos),
                selectinload(models.Item.tags),
                selectinload(models.Item.links),
                selectinload(models.Item.price_points),
            )
            .order_by(models.Item.name)
        ).all()
    )


def _path(by_id: dict, class_id: int | None) -> str | None:
    return " › ".join(crud.class_path(by_id, class_id) or []) or None


def _tags(item: models.Item) -> str | None:
    return ", ".join(sorted(tag.name for tag in item.tags)) or None


def _estimated_base(item: models.Item, base: str, rates: dict[str, float]) -> float | None:
    if item.estimated_value is None:
        return None
    return round(
        currency.to_base(float(item.estimated_value), item.currency, base, rates), 2
    )


def _purchase_base(item: models.Item, base: str, rates: dict[str, float]) -> float | None:
    if item.purchase_eur_value is not None:
        return round(float(item.purchase_eur_value) * rates.get(base, 1.0), 2)
    if item.purchase_price is None:
        return None
    return round(currency.to_base(float(item.purchase_price), item.currency, base, rates), 2)


def _gain_base(item: models.Item, base: str, rates: dict[str, float]) -> float | None:
    estimated = _estimated_base(item, base, rates)
    purchase = _purchase_base(item, base, rates)
    if estimated is None or purchase is None:
        return None
    return round(estimated - purchase, 2)


def _last_point(item: models.Item) -> models.PricePoint | None:
    return item.price_points[-1] if item.price_points else None


def _point_base(point, base: str, rates: dict[str, float]) -> float | None:
    if point is None:
        return None
    value = point.eur_value
    if value is None:
        value = currency.to_eur(float(point.price), point.currency, rates)
    return round(currency.to_base(float(value), "EUR", base, rates), 2)


# --------------------------------------------------------------------------- #
# Mise en forme
# --------------------------------------------------------------------------- #
def _write_table(ws, columns: list[Column], rows: list[models.Item], start_row: int = 1) -> int:
    for index, (label, _getter, _fmt) in enumerate(columns, start=1):
        cell = ws.cell(row=start_row, column=index, value=label)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for offset, item in enumerate(rows, start=1):
        for index, (_label, getter, fmt) in enumerate(columns, start=1):
            cell = ws.cell(row=start_row + offset, column=index, value=getter(item))
            if fmt == "money":
                cell.number_format = MONEY_FORMAT
            elif fmt == "date":
                cell.number_format = DATE_FORMAT
            elif fmt == "wrap":
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    last_row = start_row + len(rows)
    ws.freeze_panes = f"A{start_row + 1}"
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(columns))}{last_row}"
    _autosize(ws, len(columns), last_row)
    return last_row + 1


def _autosize(ws, column_count: int, last_row: int) -> None:
    for column in range(1, column_count + 1):
        width = 10
        for row in range(1, last_row + 1):
            value = ws.cell(row=row, column=column).value
            if value is not None:
                width = max(width, len(str(value)) + 2)
        ws.column_dimensions[get_column_letter(column)].width = min(45, width)


def _block(ws, start_row: int, title: str, pairs: list[tuple[str, object]], money=False) -> int:
    ws.cell(row=start_row, column=1, value=title).font = _BLOCK_FONT
    header = start_row + 1
    for column, label in ((1, "Libellé"), (2, "Valeur")):
        cell = ws.cell(row=header, column=column, value=label)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT

    row = header
    for label, value in pairs:
        row += 1
        ws.cell(row=row, column=1, value=label)
        cell = ws.cell(row=row, column=2, value=value)
        if money and isinstance(value, float):
            cell.number_format = MONEY_FORMAT
    return row + 2


# --------------------------------------------------------------------------- #
# Colonnes
# --------------------------------------------------------------------------- #
def _collection_columns(by_id: dict, base: str, rates: dict[str, float]) -> list[Column]:
    return [
        ("ID", lambda it: it.id, None),
        ("Nom", lambda it: it.name, None),
        ("Type", lambda it: TYPE_LABEL.get(it.item_type.value, it.item_type.value), None),
        ("Statut", lambda it: STATUS_LABEL.get(it.status.value, it.status.value), None),
        ("Pays", lambda it: it.country, None),
        ("Modèle", lambda it: it.aircraft_model, None),
        ("Classe", lambda it: _path(by_id, it.class_id), None),
        ("Emplacement", lambda it: it.storage_location, None),
        ("Tags", _tags, "wrap"),
        ("Date d'achat", lambda it: it.purchase_date, "date"),
        ("Prix d'achat", lambda it: it.purchase_price, "money"),
        ("Devise", lambda it: it.currency, None),
        ("Équivalent achat (EUR)", lambda it: it.purchase_eur_value, "money"),
        ("Valeur estimée", lambda it: it.estimated_value, "money"),
        (f"Valeur estimée ({base})", lambda it: _estimated_base(it, base, rates), "money"),
        (f"Plus-value ({base})", lambda it: _gain_base(it, base, rates), "money"),
        ("Échelle", lambda it: it.diecast.scale if it.diecast else None, None),
        ("Fabricant", lambda it: it.diecast.manufacturer if it.diecast else None, None),
        ("Référence", lambda it: it.diecast.part_number if it.diecast else None, None),
        (
            "Immatriculation",
            lambda it: it.diecast.registration_number if it.diecast else None,
            None,
        ),
        ("Configuration", lambda it: it.diecast.configuration if it.diecast else None, "wrap"),
        ("Unité / escadron", lambda it: it.patch.unit_squadron if it.patch else None, None),
        ("Base", lambda it: it.patch.base if it.patch else None, None),
        (
            "Matière",
            lambda it: MATERIAL_LABEL.get(it.patch.material.value)
            if it.patch and it.patch.material
            else None,
            None,
        ),
        (
            "Fixation",
            lambda it: ATTACHMENT_LABEL.get(it.patch.attachment_type.value)
            if it.patch and it.patch.attachment_type
            else None,
            None,
        ),
        (
            "Forme",
            lambda it: SHAPE_LABEL.get(it.patch.shape.value)
            if it.patch and it.patch.shape
            else None,
            None,
        ),
        ("Événement / mission", lambda it: it.patch.event_or_mission if it.patch else None, None),
        ("Prêté à", lambda it: it.loaned_to, None),
        ("Prêt depuis", lambda it: it.loan_date, "date"),
        ("Retour prévu", lambda it: it.loan_due_date, "date"),
        ("Photos", lambda it: len(it.photos), None),
        ("Notes", lambda it: it.notes, "wrap"),
        ("Créé le", lambda it: it.created_at.date() if it.created_at else None, "date"),
        ("Modifié le", lambda it: it.updated_at.date() if it.updated_at else None, "date"),
    ]


def _wish_columns(by_id: dict, base: str, rates: dict[str, float]) -> list[Column]:
    def gap(item: models.Item) -> float | None:
        estimated = _estimated_base(item, base, rates)
        last = _point_base(_last_point(item), base, rates)
        if estimated is None or last is None:
            return None
        return round(estimated - last, 2)

    return [
        ("ID", lambda it: it.id, None),
        ("Nom", lambda it: it.name, None),
        ("Type", lambda it: TYPE_LABEL.get(it.item_type.value, it.item_type.value), None),
        ("Pays", lambda it: it.country, None),
        ("Modèle", lambda it: it.aircraft_model, None),
        ("Classe", lambda it: _path(by_id, it.class_id), None),
        ("Budget visé", lambda it: it.estimated_value, "money"),
        ("Devise", lambda it: it.currency, None),
        (f"Budget ({base})", lambda it: _estimated_base(it, base, rates), "money"),
        ("Dernier relevé", lambda it: _last_point(it).price if _last_point(it) else None, "money"),
        ("Devise relevé", lambda it: _last_point(it).currency if _last_point(it) else None, None),
        (
            "Date du relevé",
            lambda it: _last_point(it).recorded_on if _last_point(it) else None,
            "date",
        ),
        (f"Écart ({base})", gap, "money"),
        (
            "Liens",
            lambda it: "\n".join(
                f"{link.label} — {link.url}" if link.label else link.url for link in it.links
            )
            or None,
            "wrap",
        ),
        ("Tags", _tags, "wrap"),
        ("Notes", lambda it: it.notes, "wrap"),
    ]


def _preorder_columns(by_id: dict, base: str, rates: dict[str, float]) -> list[Column]:
    return [
        ("ID", lambda it: it.id, None),
        ("Nom", lambda it: it.name, None),
        ("Type", lambda it: TYPE_LABEL.get(it.item_type.value, it.item_type.value), None),
        ("Pays", lambda it: it.country, None),
        ("Modèle", lambda it: it.aircraft_model, None),
        ("Classe", lambda it: _path(by_id, it.class_id), None),
        ("Sortie prévue", lambda it: it.preorder_release_date, "date"),
        ("Boutique / fournisseur", lambda it: it.preorder_shop, None),
        ("Acompte versé", lambda it: it.preorder_deposit, "money"),
        ("Devise", lambda it: it.currency, None),
        ("Valeur estimée", lambda it: it.estimated_value, "money"),
        (f"Valeur estimée ({base})", lambda it: _estimated_base(it, base, rates), "money"),
        ("Tags", _tags, "wrap"),
        ("Notes", lambda it: it.notes, "wrap"),
    ]


# --------------------------------------------------------------------------- #
# Synthèse
# --------------------------------------------------------------------------- #
def _write_summary(ws, stats) -> None:
    ws["A1"] = "Synthèse de la collection"
    ws["A1"].font = _TITLE_FONT
    ws["A2"] = f"Devise de référence : {stats.base_currency}"
    ws["A2"].font = _NOTE_FONT

    counts = {bucket.key: bucket.count for bucket in stats.by_status}
    owned = sum(counts.get(key, 0) for key in ("EXHIBITED", "STORED_BOX", "TRANSIT"))
    gain = stats.total_estimated_value - stats.total_purchase_value

    indicators = [
        ("Objets (total)", stats.total_items),
        ("Possédés", owned),
        ("Recherchés", counts.get("WANTED", 0)),
        ("Précommandes", counts.get("PREORDER", 0)),
        (f"Valeur estimée ({stats.base_currency})", stats.total_estimated_value),
        (f"Total des achats ({stats.base_currency})", stats.total_purchase_value),
        (f"Plus-value ({stats.base_currency})", round(gain, 2)),
        ("Photos", stats.total_photos),
        ("Taux du", (stats.rates_updated_at or "—")[:10]),
    ]
    next_row = _block(ws, 4, "Indicateurs", indicators, money=True)

    next_row = _block(
        ws,
        next_row,
        "Répartition par statut",
        [(STATUS_LABEL.get(key, key), count) for key, count in counts.items()],
    )
    next_row = _block(
        ws,
        next_row,
        "Répartition par type",
        [(TYPE_LABEL.get(b.key, b.key), b.count) for b in stats.by_type],
    )
    next_row = _block(
        ws,
        next_row,
        f"Valeur par classe ({stats.base_currency})",
        [(b.key, b.value) for b in stats.value_by_class],
        money=True,
    )
    next_row = _block(
        ws,
        next_row,
        f"Valeur par emplacement ({stats.base_currency})",
        [(b.key, b.value) for b in stats.value_by_location],
        money=True,
    )
    _block(
        ws,
        next_row,
        "Objets par pays",
        [(b.key, b.count) for b in stats.by_country],
    )

    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 20


# --------------------------------------------------------------------------- #
# Classeur
# --------------------------------------------------------------------------- #
def build_export(db: Session) -> bytes:
    """Construit le classeur Excel et renvoie son contenu binaire."""
    stats = compute_stats(db, include_details=True)
    base = stats.base_currency
    rates = currency.get_rates(db)
    by_id = {node.id: node for node in crud.all_classes(db)}
    items = _items(db)

    owned = [
        item
        for item in items
        if item.status not in (models.ItemStatus.WANTED, models.ItemStatus.PREORDER)
    ]
    wanted = [item for item in items if item.status == models.ItemStatus.WANTED]
    preorders = [item for item in items if item.status == models.ItemStatus.PREORDER]

    workbook = Workbook()
    collection = workbook.active
    collection.title = "Collection"
    _write_table(collection, _collection_columns(by_id, base, rates), owned)

    wishes = workbook.create_sheet("Souhaits")
    _write_table(wishes, _wish_columns(by_id, base, rates), wanted)

    preorders_sheet = workbook.create_sheet("Précommandes")
    _write_table(preorders_sheet, _preorder_columns(by_id, base, rates), preorders)

    summary = workbook.create_sheet("Synthèse")
    _write_summary(summary, stats)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
